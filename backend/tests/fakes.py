"""Deterministic stand-ins for the services: a bag-of-words embedder, real PDF/DOCX builders, and
a scripted Claude that drives the *real* SDK tool runner, tools, repository and guardrails.

The scripted client is a real `anthropic.Anthropic` whose `beta.messages.parse` / `.create`
are replaced. Agent turns are scripted as `ToolCalls` / `Answer` / `REFUSAL`; verifier calls
return scripted verdicts. Citations in an `Answer` refer to the search_result blocks the real
tools returned, by index, so a test never has to know chunk ids in advance.
"""

import io
import math
import re
import zlib
from collections.abc import Callable
from dataclasses import dataclass, field
from types import SimpleNamespace
from typing import Any

import anthropic
from anthropic.types.beta import BetaMessage
from docx import Document

from app.errors import AppError
from app.grounding.verifier import Verification

# ── Files ────────────────────────────────────────────────────────────────────


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page (empty string = blank page)."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    objects.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(pages))
    for index, text in enumerate(pages):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 11 Tf 50 750 Td ({escaped}) Tj ET".encode("latin-1") if text else b""
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
            b" /Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
            % (font_id, page_ids[index] + 1)
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    out.writelines(b"%010d 00000 n \n" % offset for offset in offsets)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    """blocks: ("h1".."h3", text) for headings, ("p", text) for paragraphs,
    ("table", "a|b;c|d") for a table with rows separated by ';' and cells by '|'."""
    document = Document()
    for kind, text in blocks:
        if kind.startswith("h"):
            document.add_heading(text, level=int(kind[1]))
        elif kind == "table":
            rows = [row.split("|") for row in text.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, cells in enumerate(rows):
                for c, value in enumerate(cells):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── Embeddings ───────────────────────────────────────────────────────────────


class FakeEmbedder:
    """Bag-of-words hashing into a fixed dimension: texts sharing words get similar vectors,
    which is enough to exercise retrieval without running Ollama."""

    model_id = "fake/bag-of-words"

    def __init__(self, dimension: int = 768) -> None:
        self.dimension = dimension
        self.calls = 0

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.calls += 1
        return [self._vector(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        self.calls += 1
        return self._vector(text)

    def check(self) -> str | None:
        return None

    def _vector(self, text: str) -> list[float]:
        vector = [0.0] * self.dimension
        for word in re.findall(r"[a-z0-9]+", text.lower()):
            vector[zlib.crc32(word.encode()) % self.dimension] += 1.0
        norm = math.sqrt(sum(v * v for v in vector)) or 1.0
        return [v / norm for v in vector]


class UnavailableEmbedder(FakeEmbedder):
    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        raise AppError("EMBEDDING_UNAVAILABLE", "The embedding service is unavailable.", 503)

    def embed_query(self, text: str) -> list[float]:
        raise AppError("EMBEDDING_UNAVAILABLE", "The embedding service is unavailable.", 503)

    def check(self) -> str | None:
        return "Ollama isn't reachable."


# ── Scripted Claude ──────────────────────────────────────────────────────────

REFUSAL = "refusal"

# A citation spec inside an Answer part:
#   3            → cite search_result block #3 (0-based, in request order; -1 = last), first unit
#   (3, 1, 3)    → cite its units [1:3]
#   {"document": "policy.md", "start": 0, "end": 1, "cited_text": "forged"}
#                → cite the first result from that document (cited_text override optional)
#   {"source": "kb://bogus", "cited_text": "x"} → an arbitrary citation (for validation tests)
CiteSpec = int | tuple[int, int, int] | dict[str, Any]


@dataclass(frozen=True)
class ToolCalls:
    """An assistant turn that calls tools: [(tool_name, input), ...]."""

    calls: list[tuple[str, dict[str, Any]]]


@dataclass(frozen=True)
class Answer:
    """An assistant turn with text: [(text, [cite specs]), ...]."""

    parts: list[tuple[str, list[CiteSpec]]]


Turn = ToolCalls | Answer | str | Exception | Callable[[dict[str, Any]], Any]


@dataclass
class Script:
    turns: list[Turn]
    verdicts: list[list[str]] = field(default_factory=list)
    agent_requests: list[dict[str, Any]] = field(default_factory=list)
    create_requests: list[dict[str, Any]] = field(default_factory=list)
    verify_requests: list[dict[str, Any]] = field(default_factory=list)
    _turn_index: int = 0
    _ids: int = 0

    # The SDK runner calls `parse` for agent turns; the verifier calls it with `output_format`.
    def parse(self, **request: Any) -> Any:
        if isinstance(request.get("output_format"), type):  # the SDK runner passes a sentinel
            return self._verify(request)
        self.agent_requests.append(request)
        return self._next_turn(request)

    # The wrap-up and retry calls use `create`.
    def create(self, **request: Any) -> Any:
        self.create_requests.append(request)
        return self._next_turn(request)

    def _next_turn(self, request: dict[str, Any]) -> BetaMessage:
        assert self.turns, "Claude was called but the test scripted no turns"
        turn = self.turns[min(self._turn_index, len(self.turns) - 1)]
        self._turn_index += 1
        if callable(turn) and not isinstance(turn, Exception):
            turn = turn(request)
        if isinstance(turn, Exception):
            raise turn
        if turn == REFUSAL:
            return self._message([], "refusal")
        if isinstance(turn, ToolCalls):
            blocks = [
                {"type": "tool_use", "id": self._next_id("toolu"), "name": name, "input": inp}
                for name, inp in turn.calls
            ]
            return self._message(blocks, "tool_use")
        assert isinstance(turn, Answer), f"unexpected scripted turn: {turn!r}"
        results = _search_results(request["messages"])
        blocks = []
        for text, cites in turn.parts:
            citations = [_citation(spec, results) for spec in cites]
            blocks.append({"type": "text", "text": text, "citations": citations or None})
        return self._message(blocks, "end_turn")

    def _verify(self, request: dict[str, Any]) -> Any:
        self.verify_requests.append(request)
        count = len(_numbered_sentences(request))
        if self.verdicts:
            verdicts = self.verdicts[min(len(self.verify_requests), len(self.verdicts)) - 1]
        else:
            verdicts = ["supported"] * count
        parsed = Verification.model_validate(
            {
                "verdicts": [
                    {"sentence": n, "verdict": v, "reason": ""}
                    for n, v in enumerate(verdicts, start=1)
                ]
            }
        )
        return SimpleNamespace(stop_reason="end_turn", parsed_output=parsed)

    def _message(self, content: list[dict[str, Any]], stop_reason: str) -> BetaMessage:
        return BetaMessage.model_validate(
            {
                "id": self._next_id("msg"),
                "type": "message",
                "role": "assistant",
                "model": "scripted",
                "content": content,
                "stop_reason": stop_reason,
                "stop_sequence": None,
                "usage": {"input_tokens": 1, "output_tokens": 1},
            }
        )

    def _next_id(self, prefix: str) -> str:
        self._ids += 1
        return f"{prefix}_{self._ids:04d}"


def scripted_claude(turns: list[Turn], verdicts: list[list[str]] = ()) -> anthropic.Anthropic:
    """A real SDK client whose API calls are answered by the script (no network)."""
    client = anthropic.Anthropic(api_key="test-key", max_retries=0)
    script = Script(list(turns), [list(v) for v in verdicts])
    client.beta.messages.parse = script.parse  # type: ignore[method-assign]
    client.beta.messages.create = script.create  # type: ignore[method-assign]
    client.models.retrieve = lambda model: SimpleNamespace(id=model)  # type: ignore[method-assign]
    client.with_options = lambda **_options: client  # type: ignore[method-assign]
    client.script = script  # type: ignore[attr-defined]
    return client


def _search_results(messages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Every search_result block in the conversation, in the order the API counts them."""
    found: list[dict[str, Any]] = []
    for message in messages:
        content = message.get("content") if isinstance(message, dict) else message.content
        if isinstance(content, str):
            continue
        for block in content:
            block = block if isinstance(block, dict) else block.model_dump()
            if block.get("type") == "search_result":
                found.append(block)
            elif block.get("type") == "tool_result" and not isinstance(block.get("content"), str):
                for inner in block.get("content") or []:
                    if inner.get("type") == "search_result":
                        found.append(inner)
    return found


def _citation(spec: CiteSpec, results: list[dict[str, Any]]) -> dict[str, Any]:
    if isinstance(spec, dict) and "document" in spec:
        # Cite the first result from a given document (by title prefix), optionally one unit
        # range, optionally with a forged cited_text (to exercise validation).
        matches = [r for r in results if (r.get("title") or "").startswith(spec["document"])]
        assert matches, f"no search_result from {spec['document']!r} in the conversation"
        index = results.index(matches[0])
        start, end = spec.get("start", 0), spec.get("end", 1)
        citation = _citation((index, start, end), results)
        if "cited_text" in spec:
            citation["cited_text"] = spec["cited_text"]
        return citation
    if isinstance(spec, dict):
        return {
            "type": "search_result_location",
            "source": spec["source"],
            "title": spec.get("title", "t"),
            "search_result_index": spec.get("index", 0),
            "start_block_index": spec.get("start", 0),
            "end_block_index": spec.get("end", 1),
            "cited_text": spec["cited_text"],
        }
    index, start, end = (spec, 0, 1) if isinstance(spec, int) else spec
    assert -len(results) <= index < len(results), (
        f"the script cites result #{index} but only {len(results)} exist"
    )
    block = results[index]
    units = [unit["text"] for unit in block["content"]]
    return {
        "type": "search_result_location",
        "source": block["source"],
        "title": block.get("title"),
        "search_result_index": index,
        "start_block_index": start,
        "end_block_index": end,
        "cited_text": "".join(units[start:end]),
    }


def _numbered_sentences(request: dict[str, Any]) -> list[str]:
    prompt = request["messages"][0]["content"]
    body = prompt.split("<answer_sentences>", 1)[1].split("</answer_sentences>", 1)[0]
    return [line for line in body.strip().splitlines() if re.match(r"^\d+\. ", line)]
