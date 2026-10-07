"""Builds the binary sample documents (PDF, DOCX) next to this script.

The text-based samples (CSV, Markdown, JSON, TXT) are committed as-is. Run from the project root:

    uv run --project backend python sample-documents/generate.py

Everything here is fictional: Lakeside Health Partners does not exist.
"""

import io
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent

PDF_PAGES = [
    "Lakeside Health Partners - Quality Targets 2026. This document sets the quality targets "
    "that every clinical department of Lakeside Health Partners reports against each quarter.",
    "Readmission target. For the third quarter of 2026 the 30-day readmission target is 11.5 "
    "percent for every department. A department is above target when its monthly rate exceeds "
    "11.5 percent. The patient satisfaction target is 4.3 out of 5.",
    "Escalation. A department above the readmission target for two consecutive months must "
    "present an improvement plan to the Quality Committee within 30 days. The Quality Committee "
    "is chaired by the Chief Medical Officer.",
]

DOCX_BLOCKS = [
    ("h1", "Lakeside Health Partners - Employee Benefits Summary 2026"),
    ("p", "This summary describes the benefits available to permanent employees. It is fictional."),
    ("h2", "Leave entitlements"),
    ("table", "Leave type|Days per year|Notes;Annual leave|25|Pro-rated for part-time staff;Sick leave|10|Full pay;Parental leave|16 weeks|Primary caregiver;Study leave|5|Requires manager approval"),
    ("h2", "Health cover"),
    ("p", "All permanent employees are enrolled in the Lakeside Care plan from their first day. Dependants can be added within 60 days of the start date."),
    ("h2", "Retirement"),
    ("p", "Lakeside matches employee pension contributions up to 6 percent of base salary."),
]


def make_pdf(pages: list[str]) -> bytes:
    """A minimal, valid PDF with one paragraph of Helvetica text per page (wrapped by hand)."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    objects.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(pages))
    for index, text in enumerate(pages):
        lines = _wrap(text, 90)
        ops = ["BT /F1 11 Tf 50 750 Td 14 TL"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({escaped}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
            b" /Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
            % (font_id, page_ids[index] + 1)
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    out.writelines(b"%010d 00000 n \n" % offset for offset in offsets)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def _wrap(text: str, width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if len(candidate) > width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    document = Document()
    for kind, text in blocks:
        if kind.startswith("h"):
            document.add_heading(text, level=int(kind[1]))
        elif kind == "table":
            rows = [row.split("|") for row in text.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, value in enumerate(cells):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    (HERE / "lakeside-quality-targets-2026.pdf").write_bytes(make_pdf(PDF_PAGES))
    (HERE / "lakeside-benefits-summary-2026.docx").write_bytes(make_docx(DOCX_BLOCKS))
    print("wrote lakeside-quality-targets-2026.pdf and lakeside-benefits-summary-2026.docx")
